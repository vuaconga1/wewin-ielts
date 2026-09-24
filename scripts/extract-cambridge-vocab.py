# -*- coding: utf-8 -*-
"""Extract Cambridge Vocabulary for IELTS docx files into JSON seed data."""
from __future__ import annotations

import json
import os
import re

from docx import Document

BASE = r"e:\Wewin\Wewin-Education-main\anh_wewin\flyer\unit 7\Castle & Environment"
OUT_PATH = r"e:\Wewin\Wewin-ielts\src\lib\learn\data\cambridge-vocab-units.json"

SECTION_RE = re.compile(
    r"^(NOUNS|VERBS|ADJECTIVES|ADVERBS|PHRASES|IDIOMS|COLL|OTHER|PHRASES WITH)",
    re.I,
)
UNIT_RE = re.compile(r"^UNIT\s+(\d+)\s*:\s*(.+)$", re.I)
EX_RE = re.compile(r"^Ex\s*:\s*(.+)$", re.I)
IPA_RE = re.compile(r"/[^/]+/")

TOPIC_META = {
    1: (
        "Growing Up",
        "Growing Up",
        "Family, childhood, relationships.",
        "Gia đình, tuổi thơ, mối quan hệ.",
    ),
    2: (
        "Mental and Physical Development",
        "Mental and Physical Development",
        "Mind, growth, abilities.",
        "Trí tuệ, tăng trưởng, năng lực.",
    ),
    3: (
        "Keeping Fit",
        "Keeping Fit",
        "Fitness, diet, exercise.",
        "Thể lực, dinh dưỡng, tập luyện.",
    ),
    4: (
        "Lifestyles",
        "Lifestyles",
        "Living habits and daily life.",
        "Thói quen sống và đời sống hàng ngày.",
    ),
    5: (
        "Student Life",
        "Student Life",
        "College, study, campus life.",
        "Đại học, học tập, đời sống khuôn viên.",
    ),
    6: (
        "Effective Communication",
        "Effective Communication",
        "Speaking, listening, idioms.",
        "Nói, nghe và thành ngữ giao tiếp.",
    ),
    7: (
        "On the Move",
        "On the Move",
        "Travel and transport vocabulary.",
        "Từ vựng du lịch và phương tiện.",
    ),
    8: (
        "Through the Ages",
        "Through the Ages",
        "History and time phrases.",
        "Lịch sử và cụm từ về thời gian.",
    ),
}


def clean(s: str) -> str:
    if not s:
        return ""
    s = s.replace("\xa0", " ").replace("\u200b", "")
    return re.sub(r"\s+", " ", s).strip()


def split_word_pos(raw: str) -> tuple[str, str | None]:
    raw = clean(raw)
    m = re.match(r"^(.+?)\s*\(([^)]+)\)\s*$", raw)
    if m:
        return clean(m.group(1)), clean(m.group(2)).lower()
    return raw, None


def pos_from_section(sec: str) -> str | None:
    s = sec.upper()
    if s.startswith("NOUN"):
        return "n"
    if s.startswith("VERB"):
        return "v"
    if s.startswith("ADJ"):
        return "adj"
    if s.startswith("ADV"):
        return "adv"
    if "IDIOM" in s:
        return "idiom"
    if "PHRASE" in s:
        return "phrase"
    return None


def parse_word_line(line: str, default_pos: str | None) -> dict | None:
    line = clean(line)
    if not line or SECTION_RE.match(line) or UNIT_RE.match(line):
        return None
    if EX_RE.match(line):
        return None

    ipa_m = IPA_RE.search(line)
    ipa = ipa_m.group(0) if ipa_m else None

    if ipa_m:
        before = clean(line[: ipa_m.start()].rstrip(" :"))
        after = clean(line[ipa_m.end() :].lstrip(" :"))
    elif ":" in line:
        before, after = line.split(":", 1)
        before, after = clean(before), clean(after)
    else:
        return None

    word, pos = split_word_pos(before)
    word = word.rstrip(" :").strip()
    if not word or len(word) > 80:
        return None

    meaning_vi = ""
    definition_en = ""
    if after:
        if ":" in after:
            left, right = after.rsplit(":", 1)
            left, right = clean(left), clean(right)
            if left.startswith("="):
                definition_en = left
                meaning_vi = right
            elif re.search(r"[A-Za-z]", left) and re.search(
                r"[\u00C0-\u1EF9a-zA-ZàáảãạăằắẳẵặâầấẩẫậèéẻẽẹêềếểễệìíỉĩịòóỏõọôồốổỗộơờớởỡợùúủũụưừứửữựỳýỷỹỵđĐ]",
                right,
            ):
                definition_en = left.lstrip("= ").strip()
                meaning_vi = right
            else:
                meaning_vi = right or left
                if not re.search(r"[\u00C0-\u1EF9àáảãạ]", meaning_vi) and re.search(
                    r"[A-Za-z]{3,}", meaning_vi
                ):
                    definition_en = meaning_vi
                    meaning_vi = ""
        else:
            if re.search(r"[\u00C0-\u1EF9àáảãạăâêôơưđ]", after):
                meaning_vi = after
            else:
                definition_en = after.lstrip("= ").strip()

    if not meaning_vi and not definition_en:
        return None
    if not meaning_vi:
        meaning_vi = definition_en or word

    return {
        "word": word,
        "pos": pos or default_pos,
        "ipa": ipa,
        "meaningVi": meaning_vi,
        "definitionEn": definition_en or None,
        "exampleEn": "",
        "exampleVi": None,
    }


def wrap_example(word: str, example: str) -> str:
    if not example or "[" in example or not word:
        return example
    w0 = word.split()[0]
    if w0.lower() not in example.lower():
        return example
    return re.sub(re.escape(w0), lambda m: f"[{m.group(0)}]", example, count=1, flags=re.I)


def parse_unit_1_8(path: str) -> dict[tuple[int, str], list[dict]]:
    d = Document(path)
    paras = [clean(p.text) for p in d.paragraphs]
    units: dict[tuple[int, str], list[dict]] = {}
    current_unit: tuple[int, str] | None = None
    current_sec = "NOUNS"
    pending: dict | None = None

    def flush() -> None:
        nonlocal pending
        if pending and current_unit:
            units.setdefault(current_unit, []).append(pending)
        pending = None

    for p in paras:
        if not p:
            continue
        um = UNIT_RE.match(p)
        if um:
            flush()
            current_unit = (int(um.group(1)), clean(um.group(2)).title())
            current_sec = "NOUNS"
            continue
        if SECTION_RE.match(p):
            flush()
            current_sec = p.rstrip(":")
            continue
        em = EX_RE.match(p)
        if em and pending:
            pending["exampleEn"] = wrap_example(pending["word"], clean(em.group(1)))
            continue
        flush()
        parsed = parse_word_line(p, pos_from_section(current_sec))
        if parsed:
            pending = parsed
    flush()
    return units


def parse_table_docx(path: str, has_vi: bool) -> tuple[str, list[dict]]:
    d = Document(path)
    paras = [clean(p.text) for p in d.paragraphs if clean(p.text)]
    title = paras[0] if paras else os.path.basename(path)
    words: list[dict] = []
    for table in d.tables:
        rows = table.rows
        if not rows:
            continue
        for row in rows[1:]:
            cells = [clean(c.text) for c in row.cells]
            if len(cells) < 3:
                continue
            raw_word, ipa, col2 = cells[0], cells[1], cells[2]
            col3 = cells[3] if len(cells) > 3 else ""
            word, pos = split_word_pos(raw_word)
            if not word or word.lower() in ("từ vựng", "word / phrase", "word"):
                continue
            if has_vi:
                definition_en = col2
                meaning_vi = col3 or col2
                example_en = f"The word [{word}] is useful in IELTS."
            else:
                definition_en = col2
                meaning_vi = col2
                example_en = wrap_example(word, col3) or f"The word [{word}] is useful in IELTS."
            ipa_out = None
            if ipa:
                ipa_out = ipa if ipa.startswith("/") else f"/{ipa}/"
            words.append(
                {
                    "word": word,
                    "pos": pos,
                    "ipa": ipa_out,
                    "meaningVi": meaning_vi,
                    "definitionEn": definition_en,
                    "exampleEn": example_en,
                    "exampleVi": None,
                }
            )
    return title, words


def clean_words(words: list[dict]) -> list[dict]:
    out: list[dict] = []
    for w in words:
        item = {k: v for k, v in w.items() if v is not None and v != ""}
        if "exampleEn" not in item:
            item["exampleEn"] = f'The word [{w["word"]}] is useful in IELTS.'
        if "meaningVi" not in item:
            item["meaningVi"] = w.get("definitionEn") or w["word"]
        out.append(item)
    return out


def slugify(title: str) -> str:
    s = title.lower().replace(" & ", "-").replace(" ", "-")
    return re.sub(r"[^a-z0-9-]", "", s)


def main() -> None:
    u18 = parse_unit_1_8(os.path.join(BASE, "UNIT 1-8.docx"))
    _, words9 = parse_table_docx(
        os.path.join(BASE, "UNIT 9 The Natural World.docx"), has_vi=True
    )
    _, words12 = parse_table_docx(
        os.path.join(
            BASE, "Vocabulary Summary - Unit 12 (Cambridge Vocabulary for IELTS).docx"
        ),
        has_vi=False,
    )
    _, words13 = parse_table_docx(
        os.path.join(
            BASE, "Vocabulary Summary - Unit 13 (Cambridge Vocabulary for IELTS).docx"
        ),
        has_vi=False,
    )

    topics: list[dict] = []
    order = 1
    for num in range(1, 9):
        key = next((k for k in u18 if k[0] == num), None)
        words = clean_words(u18.get(key, []) if key else [])
        en_title, _vi_title, en_sum, vi_sum = TOPIC_META[num]
        topics.append(
            {
                "slug": f"unit-{num}-{slugify(en_title)}",
                "order": order,
                "title": {
                    "vi": en_title,
                    "en": en_title,
                },
                "summary": {"vi": vi_sum, "en": en_sum},
                "words": words,
            }
        )
        print(f"Unit {num}: {len(words)} words")
        order += 1

    topics.append(
        {
            "slug": "unit-9-the-natural-world",
            "order": order,
            "title": {
                "vi": "The Natural World",
                "en": "The Natural World",
            },
            "summary": {
                "vi": "Flora, fauna và nông nghiệp — Cambridge Vocabulary for IELTS.",
                "en": "Flora, fauna and agriculture — Cambridge Vocabulary for IELTS.",
            },
            "words": clean_words(words9),
        }
    )
    print(f"Unit 9: {len(words9)} words")
    order += 1

    topics.append(
        {
            "slug": "unit-12-information-technology",
            "order": order,
            "title": {
                "vi": "Information Technology",
                "en": "Information Technology",
            },
            "summary": {
                "vi": "Viễn thông, máy tính và công nghệ.",
                "en": "Telecommunications, computers, and technology.",
            },
            "words": clean_words(words12),
        }
    )
    print(f"Unit 12: {len(words12)} words")
    order += 1

    topics.append(
        {
            "slug": "unit-13-the-modern-world",
            "order": order,
            "title": {
                "vi": "The Modern World",
                "en": "The Modern World",
            },
            "summary": {
                "vi": "Toàn cầu hóa, thái độ và xu hướng.",
                "en": "Globalisation, changing attitudes, and trends.",
            },
            "words": clean_words(words13),
        }
    )
    print(f"Unit 13: {len(words13)} words")

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(
            {"source": "Cambridge Vocabulary for IELTS", "topics": topics},
            f,
            ensure_ascii=False,
            indent=2,
        )

    total = sum(len(t["words"]) for t in topics)
    print(f"TOTAL topics={len(topics)} words={total}")
    print("wrote", OUT_PATH)
    print("sample u1:", json.dumps(topics[0]["words"][:2], ensure_ascii=False))
    print("sample u12:", json.dumps(topics[9]["words"][:1], ensure_ascii=False))


if __name__ == "__main__":
    main()
